"""DOCX parser for AXIOM ingestion workflows."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document

from backend.app.models.schemas import ExtractedImage, ExtractedTable, ExtractedTextBlock, ExtractionMethod, IngestionIssue, SeverityLevel
from backend.app.services.ingestion.ingestion_types import ParsedDocument


logger = logging.getLogger(__name__)


class DOCXParserError(RuntimeError):
    """Raised when a DOCX file cannot be parsed."""


class DOCXParser:
    """Extract paragraphs, tables, and embedded image metadata from DOCX files."""

    def __init__(self) -> None:
        """Initialize the parser."""
        self._logger = logger.getChild(self.__class__.__name__)

    def parse(self, file_path: str | Path, *, enable_table_extraction: bool = True) -> ParsedDocument:
        """Parse a DOCX file into normalized ingestion artifacts."""
        path = Path(file_path)
        if not path.exists():
            raise DOCXParserError(f"DOCX file not found: {path}")

        parsed = ParsedDocument(page_count=1)
        try:
            document = Document(path)
        except Exception as exc:  # pragma: no cover - depends on office file integrity
            self._logger.exception("Failed to open DOCX", extra={"path": str(path)})
            raise DOCXParserError(f"Failed to parse DOCX file: {path}") from exc

        for paragraph in document.paragraphs:
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            parsed.text_blocks.append(
                ExtractedTextBlock(
                    page_number=1,
                    text=text,
                    extraction_method=ExtractionMethod.DIRECT_TEXT,
                    confidence=0.98,
                )
            )

        if enable_table_extraction:
            parsed.tables.extend(self._extract_tables(document))

        parsed.images.extend(self._extract_images(document))

        if not parsed.text_blocks:
            parsed.issues.append(
                IngestionIssue(
                    code="docx.no_text",
                    message="No paragraphs were extracted from the DOCX document.",
                    severity=SeverityLevel.WARNING,
                    page_number=1,
                )
            )
        return parsed

    def _extract_tables(self, document: Document) -> list[ExtractedTable]:
        """Extract tables from a DOCX document."""
        tables: list[ExtractedTable] = []
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers = [header for header in rows[0] if header]
            if not headers:
                continue
            normalized_rows = [row[: len(headers)] + [""] * max(0, len(headers) - len(row)) for row in rows[1:]]
            tables.append(
                ExtractedTable(
                    page_number=1,
                    headers=headers,
                    rows=normalized_rows,
                    extraction_method=ExtractionMethod.TABLE,
                    confidence=0.93,
                )
            )
        return tables

    def _extract_images(self, document: Document) -> list[ExtractedImage]:
        """Extract embedded image metadata from a DOCX document."""
        images: list[ExtractedImage] = []
        for rel in document.part.rels.values():
            if "image" not in rel.target_ref:
                continue
            image_part = rel.target_part
            width = getattr(image_part, "default_cx", None) or 1
            height = getattr(image_part, "default_cy", None) or 1
            images.append(
                ExtractedImage(
                    page_number=1,
                    content_type=image_part.content_type,
                    width=max(int(width), 1),
                    height=max(int(height), 1),
                    caption=f"Embedded DOCX image {len(images) + 1}",
                )
            )
        return images
